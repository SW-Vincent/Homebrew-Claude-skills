# -*- coding: utf-8 -*-
"""
Excel workbook generator for tag management audits (GTM, Matomo Tag Manager,
Tag Commander, one or several tools on the same site).

The workbook language follows the user's language: set "lang" to "en" (default)
or "fr" in the spec. Cell contents (names, notes, statuses) must be written in
that same language by the caller.

CLI:
    python build_audit_xlsx.py spec.json output.xlsx

Import:
    from build_audit_xlsx import build_workbook
    build_workbook(spec_dict, "output.xlsx")

After generation ALWAYS run the recalculation (formulas have no cached value
until then):
    python /mnt/skills/public/xlsx/scripts/recalc.py output.xlsx

--------------------------------------------------------------------
Spec schema (full example: references/example_spec.json)

{
  "lang": "en" | "fr",
  "title": str, "subtitle": str, "method_note": str,
  "sources": [                        # one row per audited container (Summary sheet)
    {"platform": "GTM", "container": str, "version": str,
     "source": "Client code" | "Interface export" | "API export",
     "generated": str, "limits": str}
  ],
  "kpis": [                           # qualitative key figures, added AFTER the 4 automatic ones
    {"label": str, "value": str|int, "formula": str|None}
    # "formula" may use {ID} {PLATFORM} {CATEGORY} {STATUS}: replaced by the ranges of the
    # Tags sheet (e.g. "=COUNTIF({CATEGORY},\"Google Ads\")").
  ],
  "platforms": [str, ...],            # optional: display order (default: order of appearance in tags)
  "categories": [str, ...],           # categories / destinations, in display order
  "tags": [                           # Tags sheet, one row per tag
    {"platform": str, "id": str, "name": str, "category": str, "technical_type": str,
     "data_sent": str, "trigger": str, "blocker": str, "scope": str, "status": str, "notes": str}
  ],
  # Status vocabulary (use the words of the chosen language):
  #   en: Active | Active (scheduled) | Inactive (paused) | Inactive (expired) | Inactive (never launched)
  #   fr: Actif | Actif (planifié) | Inactif (pausé) | Inactif (expiré) | Inactif (jamais lancé)
  "consent_rows": [                   # Consent sheet: one row per mechanism and per platform
    [platform, type, source, behaviour, initialisation, update], ...
  ],
  "consent_extra_notes": [str, ...],
  "consent_asterisk_note": str,       # optional legend text (default provided). Any consent value
                                      # ending with "*" displays the legend automatically.
  "constats": [ {"title": str, "text": str, "highlight": bool} ]     # findings
}
--------------------------------------------------------------------
"""
import json
import math
import sys

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo

FONT_NAME = "Arial"
MAX_ROW = 2000  # formula ranges cover rows 2..2000 of the Tags sheet
TAGS_SHEET = "Tags"  # same name in every language: formulas refer to it

HEADER_FILL = PatternFill(start_color="1F3864", end_color="1F3864", fill_type="solid")
HEADER_FONT = Font(name=FONT_NAME, size=10, bold=True, color="FFFFFF")
TITLE_FONT = Font(name=FONT_NAME, size=14, bold=True, color="1F3864")
SUBTITLE_FONT = Font(name=FONT_NAME, size=10, italic=True, color="595959")
SECTION_FONT = Font(name=FONT_NAME, size=12, bold=True, color="1F3864")
BODY_FONT = Font(name=FONT_NAME, size=10)
BOLD_BODY_FONT = Font(name=FONT_NAME, size=10, bold=True)
WRAP_TOP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
ACTIVE_FILL = PatternFill(start_color="E4F3E4", end_color="E4F3E4", fill_type="solid")
INACTIVE_FILL = PatternFill(start_color="FCE4E4", end_color="FCE4E4", fill_type="solid")
ANOMALY_FILL = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")

# Fixed order of the Tags sheet columns (formulas refer to them by letter): key -> letter
COLUMN_KEYS = ["platform", "id", "name", "category", "technical_type", "data_sent",
               "trigger", "blocker", "scope", "status", "notes"]
COLUMN_LETTERS = dict(zip(COLUMN_KEYS, "ABCDEFGHIJK"))
RANGE = {
    "{ID}": "%s!$B$2:$B$%d" % (TAGS_SHEET, MAX_ROW),
    "{PLATFORM}": "%s!$A$2:$A$%d" % (TAGS_SHEET, MAX_ROW),
    "{CATEGORY}": "%s!$D$2:$D$%d" % (TAGS_SHEET, MAX_ROW),
    "{STATUS}": "%s!$J$2:$J$%d" % (TAGS_SHEET, MAX_ROW),
}

LABELS = {
    "en": {
        "sheet_summary": "Summary", "sheet_consent": "Consent & Findings",
        "default_title": "Tag management audit",
        "default_method": "Method: static reading of the containers (tags, triggers, blockers, scopes). "
                          "To be confirmed in the Preview/Debug mode of each tool on the live site.",
        "sources_section": "Scope of the analysis and source limits",
        "sources_headers": ["Platform", "Container", "Version", "Source", "Date", "Limits of this source"],
        "kpi_section": "Key figures",
        "kpi_total": "Total number of tags analysed (excluding infrastructure tags)",
        "kpi_active": "Number of active tags",
        "kpi_inactive": "Number of inactive tags (paused, expired, never launched)",
        "kpi_platforms": "Number of platforms analysed",
        "by_platform": "Breakdown by platform", "by_platform_headers": ["Platform", "Tags", "Active", "Inactive"],
        "total": "Total",
        "by_category": "Breakdown of tags by destination / category (all platforms)",
        "by_category_headers": ["Category / Destination", "Tags"],
        "tags_headers": ["Platform", "Tag ID", "Tag name / Destination", "Category", "Technical type",
                         "Event / Data sent", "Trigger", "Blocker(s) / Exceptions", "Scope (Tag Commander)",
                         "Status", "Notes"],
        "consent_title": "Consent: mechanism detected, per platform",
        "consent_headers": ["Platform", "Mechanism / consent type", "Source (CMP or configuration)", "Behaviour",
                            "Initialisation (Default / configuration)", "Update (Update / after choice)"],
        "consent_extra": "Additional settings:",
        "findings_title": "Findings and points of attention (static analysis)",
        "asterisk": "* Verdict based on the container alone, without a real test. Consent (and Consent Mode) is often "
                    "handled upstream by the cookie banner: confirm with a test (network requests before and after "
                    "the choice in the banner).",
        "active_prefix": "Active", "inactive_prefix": "Inactive",
        "too_many": "More than %d tags: raise MAX_ROW in the script.",
    },
    "fr": {
        "sheet_summary": "Synthèse", "sheet_consent": "Consentement & Anomalies",
        "default_title": "Audit de tag management",
        "default_method": "Méthode : lecture statique des conteneurs (tags, déclencheurs, bloqueurs, périmètres). "
                          "À confirmer en mode Aperçu/Debug de chaque outil sur le site réel.",
        "sources_section": "Périmètre de l'analyse et limites des sources",
        "sources_headers": ["Plateforme", "Conteneur", "Version", "Source", "Date", "Limites de cette source"],
        "kpi_section": "Chiffres clés",
        "kpi_total": "Nombre total de tags analysés (hors tags d'infrastructure)",
        "kpi_active": "Nombre de tags actifs",
        "kpi_inactive": "Nombre de tags inactifs (pausés, expirés, jamais lancés)",
        "kpi_platforms": "Nombre de plateformes analysées",
        "by_platform": "Répartition par plateforme", "by_platform_headers": ["Plateforme", "Tags", "Actifs", "Inactifs"],
        "total": "Total",
        "by_category": "Répartition des tags par destination / catégorie (toutes plateformes)",
        "by_category_headers": ["Catégorie / Destination", "Tags"],
        "tags_headers": ["Plateforme", "ID Tag", "Nom du tag / Destination", "Catégorie", "Type technique",
                         "Événement / Données envoyées", "Déclencheur", "Bloqueur(s) / Exceptions",
                         "Périmètre (Tag Commander)", "Statut", "Remarques particulières"],
        "consent_title": "Consentement : configuration détectée, par plateforme",
        "consent_headers": ["Plateforme", "Mécanisme / type de consentement", "Source (CMP ou configuration)",
                            "Comportement", "Initialisation (Default / configuration)", "Mise à jour (Update / après choix)"],
        "consent_extra": "Paramètres complémentaires :",
        "findings_title": "Constats et points d'attention (analyse statique)",
        "asterisk": "* Verdict établi depuis le conteneur seul, sans test réel. Le consentement (et le Consent Mode) est "
                    "souvent géré en amont par le bandeau cookies : à confirmer par un test (requêtes réseau avant et "
                    "après le choix dans le bandeau).",
        "active_prefix": "Actif", "inactive_prefix": "Inactif",
        "too_many": "Plus de %d tags : augmenter MAX_ROW dans le script.",
    },
}


def _expand(formula):
    for token, rng in RANGE.items():
        formula = formula.replace(token, rng)
    return formula


def _prefix_test(range_ref, prefix):
    """Language-independent test 'the status starts with <prefix>' (Active/Inactive)."""
    return 'LEFT(%s,%d)="%s"' % (range_ref, len(prefix), prefix)


def _ends_star(v):
    return isinstance(v, str) and v.rstrip().endswith("*")


def _asterisk_note(spec, L):
    """Legend text if a consent value (key figure or consent row) ends with '*'."""
    used = any(_ends_star(k.get("value")) for k in spec.get("kpis", [])) or \
        any(_ends_star(c) for row in spec.get("consent_rows", []) for c in row)
    return spec.get("consent_asterisk_note", L["asterisk"]) if used else None


def _legend(ws, row, text, last_col):
    c = ws.cell(row=row, column=1, value=text)
    c.font, c.alignment = Font(name=FONT_NAME, size=9, italic=True, color="7F6000"), WRAP_TOP
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=last_col)
    ws.row_dimensions[row].height = 30


def _section(ws, row, text):
    ws.cell(row=row, column=1, value=text).font = SECTION_FONT
    return row + 1


def _header(ws, row, labels):
    for col, label in enumerate(labels, start=1):
        c = ws.cell(row=row, column=col, value=label)
        c.font, c.fill, c.alignment = HEADER_FONT, HEADER_FILL, CENTER
    ws.row_dimensions[row].height = 30


def _build_summary(wb, spec, platforms, L):
    ws = wb.active
    ws.title = L["sheet_summary"]
    ws.sheet_view.showGridLines = False
    act = L["active_prefix"]
    ina = L["inactive_prefix"]

    ws["A1"] = spec.get("title", L["default_title"])
    ws["A1"].font = TITLE_FONT
    ws["A2"] = spec.get("subtitle", "")
    ws["A2"].font = SUBTITLE_FONT
    ws["A3"] = spec.get("method_note", L["default_method"])
    ws["A3"].font = SUBTITLE_FONT

    r = 5
    sources = spec.get("sources", [])
    if sources:
        r = _section(ws, r, L["sources_section"])
        _header(ws, r, L["sources_headers"])
        r += 1
        for s in sources:
            vals = [s.get("platform"), s.get("container"), s.get("version"), s.get("source"), s.get("generated"), s.get("limits")]
            for col, v in enumerate(vals, start=1):
                c = ws.cell(row=r, column=col, value=v)
                c.font, c.border, c.alignment = BODY_FONT, BORDER, WRAP_TOP
            ws.row_dimensions[r].height = max(30, 14 * math.ceil(len(str(s.get("limits", ""))) / 55))
            r += 1
        r += 1

    r = _section(ws, r, L["kpi_section"])
    auto = [
        (L["kpi_total"], "=COUNTA(%s)" % RANGE["{ID}"]),
        (L["kpi_active"], "=SUMPRODUCT(--(%s))" % _prefix_test(RANGE["{STATUS}"], act)),
        (L["kpi_inactive"], "=SUMPRODUCT(--(%s))" % _prefix_test(RANGE["{STATUS}"], ina)),
        (L["kpi_platforms"], len(platforms)),
    ]
    for label, val in auto:
        ws.cell(row=r, column=1, value=label).font = BODY_FONT
        c = ws.cell(row=r, column=2, value=val)
        c.font, c.alignment = BOLD_BODY_FONT, Alignment(horizontal="left")
        r += 1
    for kpi in spec.get("kpis", []):
        ws.cell(row=r, column=1, value=kpi["label"]).font = BODY_FONT
        val = _expand(kpi["formula"]) if kpi.get("formula") else kpi.get("value")
        c = ws.cell(row=r, column=2, value=val)
        c.font, c.alignment = BOLD_BODY_FONT, Alignment(horizontal="left")
        r += 1
    note = _asterisk_note(spec, L)
    if note:
        _legend(ws, r, note, 6)
        r += 1

    r += 1
    r = _section(ws, r, L["by_platform"])
    _header(ws, r, L["by_platform_headers"])
    r += 1
    first = r
    for p in platforms:
        ws.cell(row=r, column=1, value=p)
        ws.cell(row=r, column=2, value="=COUNTIF(%s,A%d)" % (RANGE["{PLATFORM}"], r))
        ws.cell(row=r, column=3, value="=SUMPRODUCT((%s=A%d)*(%s))" % (RANGE["{PLATFORM}"], r, _prefix_test(RANGE["{STATUS}"], act)))
        ws.cell(row=r, column=4, value="=SUMPRODUCT((%s=A%d)*(%s))" % (RANGE["{PLATFORM}"], r, _prefix_test(RANGE["{STATUS}"], ina)))
        for col in range(1, 5):
            c = ws.cell(row=r, column=col)
            c.font, c.border = BODY_FONT, BORDER
            if col > 1:
                c.alignment = CENTER
        r += 1
    ws.cell(row=r, column=1, value=L["total"]).font = BOLD_BODY_FONT
    for col, letter in ((2, "B"), (3, "C"), (4, "D")):
        c = ws.cell(row=r, column=col, value="=SUM(%s%d:%s%d)" % (letter, first, letter, r - 1))
        c.font, c.alignment = BOLD_BODY_FONT, CENTER

    r += 2
    r = _section(ws, r, L["by_category"])
    _header(ws, r, L["by_category_headers"])
    r += 1
    first = r
    for cat in spec.get("categories", []):
        ws.cell(row=r, column=1, value=cat)
        ws.cell(row=r, column=2, value="=COUNTIF(%s,A%d)" % (RANGE["{CATEGORY}"], r))
        for col in (1, 2):
            c = ws.cell(row=r, column=col)
            c.font, c.border = BODY_FONT, BORDER
            if col == 2:
                c.alignment = CENTER
        r += 1
    ws.cell(row=r, column=1, value=L["total"]).font = BOLD_BODY_FONT
    tot = ws.cell(row=r, column=2, value="=SUM(B%d:B%d)" % (first, r - 1))
    tot.font, tot.alignment = BOLD_BODY_FONT, CENTER

    for col, w in {"A": 58, "B": 26, "C": 16, "D": 22, "E": 16, "F": 60}.items():
        ws.column_dimensions[col].width = w
    return ws


def _build_tags(wb, spec, L):
    ws = wb.create_sheet(TAGS_SHEET)
    ws.sheet_view.showGridLines = False
    for col, label in enumerate(L["tags_headers"], start=1):
        c = ws.cell(row=1, column=col, value=label)
        c.font, c.fill, c.alignment = HEADER_FONT, HEADER_FILL, CENTER
    ws.row_dimensions[1].height = 32

    tags = spec.get("tags", [])
    if len(tags) + 1 > MAX_ROW:
        raise ValueError(L["too_many"] % (MAX_ROW - 1))
    act = L["active_prefix"].lower()
    for i, t in enumerate(tags, start=2):
        for col, key in enumerate(COLUMN_KEYS, start=1):
            cell = ws.cell(row=i, column=col, value=t.get(key, ""))
            cell.font, cell.border = BODY_FONT, BORDER
            cell.alignment = Alignment(horizontal="center", vertical="top") if key in ("platform", "id") else WRAP_TOP
        status = ws.cell(row=i, column=10)
        status.fill = ACTIVE_FILL if str(status.value).strip().lower().startswith(act) else INACTIVE_FILL
        if "ANOMAL" in str(ws.cell(row=i, column=11).value or "").upper():  # ANOMALY / ANOMALIE
            ws.cell(row=i, column=11).fill = ANOMALY_FILL

    last_row = 1 + len(tags)
    widths = {"A": 14, "B": 9, "C": 34, "D": 24, "E": 20, "F": 44, "G": 38, "H": 34, "I": 34, "J": 16, "K": 46}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "C2"
    if last_row >= 2:
        tab = Table(displayName="TagsInventory", ref="A1:K%d" % last_row)
        tab.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        ws.add_table(tab)
    return ws


def _build_consent(wb, spec, L):
    ws = wb.create_sheet(L["sheet_consent"])
    ws.sheet_view.showGridLines = False
    ws["A1"] = L["consent_title"]
    ws["A1"].font = SECTION_FONT
    _header(ws, 2, L["consent_headers"])

    r = 3
    for row_data in spec.get("consent_rows", []):
        for col, val in enumerate(row_data, start=1):
            c = ws.cell(row=r, column=col, value=val)
            c.font, c.border, c.alignment = BODY_FONT, BORDER, WRAP_TOP
        ws.row_dimensions[r].height = max(30, 14 * math.ceil(max(len(str(v)) for v in row_data) / 34))
        r += 1

    note = _asterisk_note(spec, L)
    if note:
        _legend(ws, r, note, 6)
        r += 1

    extra = spec.get("consent_extra_notes", [])
    if extra:
        r += 1
        ws.cell(row=r, column=1, value=L["consent_extra"]).font = BOLD_BODY_FONT
        r += 1
        for p in extra:
            c = ws.cell(row=r, column=1, value="•  " + p)
            c.font, c.alignment = BODY_FONT, WRAP_TOP
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
            ws.row_dimensions[r].height = max(30, 14 * math.ceil(len(p) / 120))
            r += 1

    r += 1
    ws.cell(row=r, column=1, value=L["findings_title"]).font = SECTION_FONT
    r += 1
    for item in spec.get("constats", []):
        ws.cell(row=r, column=1, value=item["title"]).font = BOLD_BODY_FONT
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
        r += 1
        cell = ws.cell(row=r, column=1, value=item["text"])
        cell.font, cell.alignment = BODY_FONT, WRAP_TOP
        if item.get("highlight"):
            for col in range(1, 7):
                ws.cell(row=r, column=col).fill = ANOMALY_FILL
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
        ws.row_dimensions[r].height = max(32, 14 * math.ceil(len(item["text"]) / 115))
        r += 2

    for col, w in {"A": 16, "B": 26, "C": 32, "D": 40, "E": 30, "F": 30}.items():
        ws.column_dimensions[col].width = w
    return ws


def build_workbook(spec: dict, output_path: str) -> str:
    """Builds the audit workbook (3 sheets) and saves it to output_path.

    Then run /mnt/skills/public/xlsx/scripts/recalc.py on the produced file.
    """
    lang = spec.get("lang", "en")
    if lang not in LABELS:
        raise ValueError("lang must be one of %s" % sorted(LABELS))
    L = LABELS[lang]
    platforms = list(spec.get("platforms") or [])
    for t in spec.get("tags", []):
        if t.get("platform") and t["platform"] not in platforms:
            platforms.append(t["platform"])
    wb = openpyxl.Workbook()
    _build_summary(wb, spec, platforms, L)
    _build_tags(wb, spec, L)
    _build_consent(wb, spec, L)
    wb.save(output_path)
    return output_path


if __name__ == "__main__":
    if len(sys.argv) != 3:
        print("Usage: python build_audit_xlsx.py spec.json output.xlsx")
        sys.exit(1)
    with open(sys.argv[1], "r", encoding="utf-8") as f:
        _spec = json.load(f)
    build_workbook(_spec, sys.argv[2])
    print("OK -> %s  (now run recalc.py on this file)" % sys.argv[2])
